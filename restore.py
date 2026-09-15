#!/usr/bin/env python3
"""Restore the room-checkout database from a backup workbook.

    python3 restore.py backups/room-checkout-2026-09-15.xlsx           # preview
    python3 restore.py backups/room-checkout-2026-09-15.xlsx --apply   # do it

Previews by default and changes nothing until you pass --apply. Restoring
makes the database match the backup exactly: rows in the backup are written,
and rows that exist now but are absent from the backup are deleted.

Writing needs an editor sign-in, so it asks for the shared department account
and its password. The password is read without echoing and is never stored or
written anywhere. The class schedule is not touched - that lives in
index.html and is restored with git, not from here.
"""
import argparse, getpass, json, pathlib, re, sys, urllib.error, urllib.request

HERE = pathlib.Path(__file__).parent
TABLES = ["reservations", "exceptions", "rooms"]
# Fields the workbook deliberately does not carry (a creation timestamp is
# noise in a human-readable sheet). They are ignored when comparing and kept
# from the existing row when writing, so restoring current state is a no-op.
CARRY = ["at"]
DAYNUM = {"sunday": 0, "monday": 1, "tuesday": 2, "wednesday": 3,
          "thursday": 4, "friday": 5, "saturday": 6}


def config():
    text = (HERE / "config.js").read_text()
    return (re.search(r'supabaseUrl:\s*"([^"]+)"', text).group(1).rstrip("/"),
            re.search(r'supabaseAnonKey:\s*"([^"]+)"', text).group(1))


def call(url, key, path, method="GET", body=None, token=None, extra=None):
    headers = {"apikey": key, "Authorization": f"Bearer {token or key}",
               "Content-Type": "application/json"}
    headers.update(extra or {})
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(f"{url}{path}", data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            raw = r.read().decode()
            return json.loads(raw) if raw.strip() else None
    except urllib.error.HTTPError as e:
        detail = e.read().decode()[:300]
        raise SystemExit(f"\n{method} {path} failed ({e.code}): {detail}")


def sign_in(url, key):
    print("Sign in as an editor to write changes.")
    email = input("  Email: ").strip()
    password = getpass.getpass("  Password (not shown): ")
    res = call(url, key, "/auth/v1/token?grant_type=password", "POST",
               {"email": email, "password": password})
    print(f"  Signed in as {res['user']['email']}\n")
    return res["access_token"]


def read_backup(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    need = {"Checkouts", "Changes", "Rooms"}
    missing = need - set(wb.sheetnames)
    if missing:
        raise SystemExit(f"{path} is missing sheet(s): {', '.join(sorted(missing))}")

    def rows(name):
        ws = wb[name]
        head = [str(c.value).strip() if c.value else "" for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None and str(v).strip() for v in r):
                yield dict(zip(head, r))

    def txt(v):
        return "" if v is None else str(v).strip()

    out = {t: {} for t in TABLES}
    for r in rows("Checkouts"):
        days = [] if txt(r["weekdays"]).lower() in ("", "every day") else [
            DAYNUM[d.strip().lower()] for d in txt(r["weekdays"]).split(";")
            if d.strip().lower() in DAYNUM]
        out["reservations"][txt(r["id"])] = {
            "title": txt(r["title"]), "kind": txt(r["category"]) or "checkout",
            "person": txt(r["person"]),
            "rooms": [x.strip() for x in txt(r["rooms"]).split(";") if x.strip()],
            "days": days, "dateStart": txt(r["date start"]), "dateEnd": txt(r["date end"]),
            "timeStart": txt(r["time start"]), "timeEnd": txt(r["time end"]),
            "note": txt(r["note"]), "cancelled": txt(r["cancelled"]).lower() in ("yes", "true")}
    for r in rows("Changes"):
        d = {"baseId": txt(r["block id"]), "date": txt(r["date"]),
             "status": txt(r["status"]), "note": txt(r["note"])}
        if txt(r["new room"]):
            d.update(room=txt(r["new room"]), start=txt(r["new start"]), end=txt(r["new end"]))
        out["exceptions"][txt(r["id"])] = d
    for r in rows("Rooms"):
        out["rooms"][txt(r["id"])] = {"name": txt(r["name"]), "number": txt(r["number"]),
                                      "label": txt(r["label"])}
    return out


def norm(d):
    """The part of a row the backup can actually represent."""
    return {k: v for k, v in d.items() if k not in CARRY}


def live(url, key):
    return {t: {r["id"]: r["data"] for r in call(url, key, f"/rest/v1/{t}?select=*")}
            for t in TABLES}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("backup", help="path to a backup .xlsx")
    ap.add_argument("--apply", action="store_true", help="actually write the changes")
    args = ap.parse_args()

    path = pathlib.Path(args.backup)
    if not path.exists():
        raise SystemExit(f"No such file: {path}")

    url, key = config()
    want, have = read_backup(path), live(url, key)

    plan = {}
    for t in TABLES:
        add = [i for i in want[t] if i not in have[t]]
        upd = [i for i in want[t] if i in have[t] and norm(have[t][i]) != norm(want[t][i])]
        rm = [i for i in have[t] if i not in want[t]]
        plan[t] = (add, upd, rm)

    label = {"reservations": "Checkouts", "exceptions": "Changes", "rooms": "Rooms"}
    total = sum(len(a) + len(u) + len(d) for a, u, d in plan.values())
    print(f"Restoring from {path.name}\n")
    for t in TABLES:
        add, upd, rm = plan[t]
        print(f"  {label[t]:10s} {len(want[t])} in backup, {len(have[t])} live "
              f"-> +{len(add)} restore, ~{len(upd)} change, -{len(rm)} delete")
        for i in rm:
            d = have[t][i]
            print(f"      delete: {d.get('title') or d.get('name') or d.get('baseId') or i}")
        for i in upd:
            d = want[t][i]
            print(f"      change: {d.get('title') or d.get('name') or d.get('baseId') or i}")
        for i in add:
            d = want[t][i]
            print(f"      restore: {d.get('title') or d.get('name') or d.get('baseId') or i}")

    if not total:
        print("\nThe database already matches this backup. Nothing to do.")
        return 0
    if not args.apply:
        print(f"\n{total} change(s). Nothing written - re-run with --apply to restore.")
        return 0

    print(f"\nThis will make the database match {path.name}, deleting anything newer.")
    if input('Type "restore" to continue: ').strip().lower() != "restore":
        print("Cancelled.")
        return 1

    token = sign_in(url, key)
    for t in TABLES:
        add, upd, rm = plan[t]
        for i in add + upd:
            keep = {k: v for k, v in have[t].get(i, {}).items() if k in CARRY}
            call(url, key, f"/rest/v1/{t}", "POST", [{"id": i, "data": {**keep, **want[t][i]}}],
                 token=token, extra={"Prefer": "resolution=merge-duplicates"})
        for i in rm:
            call(url, key, f"/rest/v1/{t}?id=eq.{urllib.parse.quote(i)}", "DELETE", token=token)
        print(f"  {label[t]}: restored {len(add)+len(upd)}, deleted {len(rm)}")
    print("\nDone. Reload the site to see it.")
    return 0


if __name__ == "__main__":
    import urllib.parse
    sys.exit(main())
